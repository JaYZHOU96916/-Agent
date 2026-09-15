"""Bounded parsing and profiling, executed by the dedicated parser worker."""

import csv
import json
import math
import zipfile
from datetime import date, datetime
from itertools import islice
from pathlib import Path

import pandas as pd
import pyarrow.parquet as pq


class InvalidDataset(ValueError):
    pass


def validate_columns(columns: list, max_columns: int) -> list[str]:
    if not columns or len(columns) > max_columns:
        raise InvalidDataset(f"Dataset must contain 1–{max_columns} columns.")
    names = [str(c) if c is not None else "" for c in columns]
    if any(not name.strip() or len(name) > 256 for name in names):
        raise InvalidDataset("Column names must be non-empty and at most 256 characters.")
    if len(set(names)) != len(names):
        raise InvalidDataset("Duplicate column names are not supported; rename them before upload.")
    return names


def read_table(path: Path, extension: str, max_rows: int, max_columns: int, max_memory: int) -> tuple[pd.DataFrame, list[str]]:
    warnings = []
    if extension == "csv":
        # UTF-8/BOM only: silent fallback decoding can corrupt column names.
        csv.field_size_limit(1024 * 1024)
        with path.open(encoding="utf-8-sig", newline="") as source:
            rows = csv.reader(source, strict=True)
            names = validate_columns(next(rows, []), max_columns)
            count = 0
            for row in rows:
                if not row:
                    continue
                count += 1
                if count > max_rows:
                    raise InvalidDataset(f"Dataset exceeds {max_rows} rows.")
                if len(row) != len(names):
                    raise InvalidDataset("CSV rows must have the same number of fields as the header.")
        frame = pd.read_csv(path, encoding="utf-8-sig", nrows=max_rows + 1)
    elif extension == "xlsx":
        from openpyxl import load_workbook

        with zipfile.ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 10_000 or sum(e.file_size for e in entries) > 100 * 1024 * 1024:
                raise InvalidDataset("Expanded Excel archive exceeds the 100 MiB limit.")
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            sheet = book.worksheets[0]
            if sheet.max_column and sheet.max_column > max_columns:
                raise InvalidDataset(f"Dataset exceeds {max_columns} columns.")
            # Dimensions can be stale; read actual rows under an explicit cap.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(values_only=True)
            names = validate_columns(list(next(rows, ())), max_columns)
            values = list(islice(rows, max_rows + 1))
            if any(len(row) > len(names) for row in values):
                raise InvalidDataset("Excel rows contain cells outside the header columns.")
            values = [tuple(row) + (None,) * (len(names) - len(row)) for row in values]
            frame = pd.DataFrame(values, columns=names)
            warnings.append("Excel: first worksheet only; formula cells use saved values and are never executed.")
        finally:
            book.close()
    elif extension == "xls":
        import xlrd

        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = book.sheet_by_index(0)
            if sheet.nrows > max_rows + 1:
                raise InvalidDataset(f"Dataset exceeds {max_rows} rows.")
            validate_columns(sheet.row_values(0) if sheet.nrows else [], max_columns)
        finally:
            book.release_resources()
        frame = pd.read_excel(path, engine="xlrd", nrows=max_rows + 1)
        warnings.append("Excel: first worksheet only; formula cells use saved values and are never executed.")
    else:
        with pq.ParquetFile(path) as parquet:
            validate_columns(parquet.schema_arrow.names, max_columns)
            if parquet.metadata.num_rows > max_rows:
                raise InvalidDataset(f"Dataset exceeds {max_rows} rows.")
            expanded = sum(parquet.metadata.row_group(i).total_byte_size for i in range(parquet.metadata.num_row_groups))
            if expanded > max_memory:
                raise InvalidDataset("Expanded Parquet data exceeds the memory budget.")
            frame = parquet.read(use_threads=False).to_pandas()
    frame.columns = validate_columns(list(frame.columns), max_columns)
    if len(frame) > max_rows:
        raise InvalidDataset(f"Dataset exceeds {max_rows} rows.")
    if int(frame.memory_usage(deep=True).sum()) > max_memory:
        raise InvalidDataset("Decoded dataset exceeds the memory budget.")
    return frame, warnings


def json_scalar(value):
    if value is None or value is pd.NA or value is pd.NaT:
        return None
    if hasattr(value, "item"):
        value = value.item()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (bool, int, float)):
        return value
    if isinstance(value, bytes):
        return "[binary data]"
    if isinstance(value, (list, dict, tuple)):
        return json.dumps(value, default=str, ensure_ascii=False)[:256]
    return str(value)[:256]


def profile_table(frame: pd.DataFrame, warnings: list[str]) -> dict:
    count = len(frame)
    columns = [
        {"name": name, "dtype": str(frame[name].dtype),
         "missing_count": int(frame[name].isna().sum()),
         "missing_fraction": float(frame[name].isna().mean()) if count else 0.0}
        for name in frame.columns
    ]
    samples = [{name: json_scalar(value) for name, value in row.items()}
               for row in frame.head(5).to_dict(orient="records")]
    return {"row_count": count, "column_count": len(columns), "columns": columns,
            "sample_rows": samples, "warnings": warnings + [
                "Preview strings are limited to 256 characters; non-finite values become null in previews."
            ]}
